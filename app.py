import os, io, re, json, uuid, zipfile, time, hashlib, logging, csv
from datetime import datetime
from typing import List, Dict, Any, Optional
from queue import Queue, Empty
from flask import Flask, request, render_template_string, send_file, redirect, url_for, Response
import requests

# ------- Optional parsers -------
try:
    from pdfminer.high_level import extract_text as pdf_extract_text
except Exception:
    pdf_extract_text = None

try:
    import docx
except Exception:
    docx = None

try:
    from bs4 import BeautifulSoup
except Exception:
    BeautifulSoup = None

# ------- Config -------
MODEL_API_KEY  = os.getenv("MODEL_API_KEY", "")
MODEL_BASE_URL = os.getenv("MODEL_BASE_URL", "https://api.openai.com")  # 可指向 DeepSeek 的 OpenAI 兼容 Base URL
MODEL_NAME     = os.getenv("MODEL_NAME", "deepseek-chat")
MAX_WORKERS    = int(os.getenv("MAX_WORKERS", "2"))  # 默认2，更稳
MAX_CHARS_EACH = int(os.getenv("MAX_CHARS_EACH", "12000"))
TIMEOUT_SEC    = int(os.getenv("TIMEOUT_SEC", "90"))
RETRIES        = int(os.getenv("RETRIES", "2"))

ALLOWED_EXT = {".pdf", ".docx", ".txt", ".csv", ".zip", ".html", ".htm"}

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 300 * 1024 * 1024  # 300MB
logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")

REPORTS: Dict[str, Dict[str, Any]] = {}
JOBS: Dict[str, Dict[str, Any]] = {}   # rid -> {"q": Queue, "done": False, "title": str}

# ------- HTML Templates -------
INDEX_HTML = """<!DOCTYPE html><html lang="zh"><head><meta charset="utf-8"/><meta name="viewport" content="width=device-width,initial-scale=1"/>
<title>Alsos Talent · 合规AI自动化寻访（MVP）</title>
<style>
 body { font-family:-apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,Arial;margin:0;background:#0b0f14;color:#e3e8f2;}
 .wrap { max-width: 980px; margin: 32px auto; padding: 0 16px; }
 h1 { font-size: 22px; margin: 12px 0 18px; }
 .card { background:#121824; border:1px solid #1e2633; border-radius:16px; padding:20px; margin-bottom:18px; }
 label { display:block; font-size:14px; color:#A9B4C6; margin:8px 0 6px; }
 input[type="text"], textarea { width:100%; background:#0b1018; color:#dbe4f0; border:1px solid #223044; border-radius:10px; padding:10px 12px; outline:none; }
 textarea { min-height: 110px; }
 .row { display:grid; grid-template-columns:1fr 1fr; gap:16px; }
 .btn { background:#2563eb; color:white; border:none; padding:12px 16px; border-radius:12px; cursor:pointer; font-weight:600; }
 small { color:#93a1b7; } .muted { color:#93a1b7; font-size:12px; } .pill { display:inline-block; padding:2px 8px; background:#102033; border:1px solid #223044; border-radius:999px; margin-right:6px; font-size:12px; color:#B8C4D9;}
 a{ color:#7aa0ff; text-decoration:none;}
</style></head><body><div class="wrap">
  <h1>Alsos Talent · 合规AI自动化寻访（MVP）</h1>
  <div class="card"><p class="muted">说明：本工具<strong>不做</strong>对 LinkedIn/猎聘 的自动点开或抓取；仅对你<strong>合规导出</strong>的 ZIP/PDF/HTML/CSV/文本做AI分析、排序并导出Excel。</p></div>
  <form action="/process" method="post" enctype="multipart/form-data">
    <div class="card"><h3>上传候选集（支持多文件）</h3>
      <label>选择文件（.zip .pdf .html/.htm .docx .txt .csv）：</label>
      <input type="file" name="files" multiple required />
      <small>可直接上传 Recruiter Lite 25人/包的 ZIP（一次多包）。</small>
    </div>
    <div class="card"><h3>岗位/筛选要求</h3>
      <div class="row">
        <div><label>职位名称 / 方向</label><input type="text" name="role" placeholder="例如：VP/SVP of Biology（免疫/肿瘤）"/></div>
        <div><label>最低年限</label><input type="text" name="min_years" placeholder="例如：8 或 10-15"/></div>
      </div>
      <div class="row">
        <div><label>Must-have关键词（逗号分隔）</label><input type="text" name="must" placeholder="例如：ADC, 临床前, 抗体工程, 跨职能领导"/></div>
        <div><label>Nice-to-have关键词（逗号分隔）</label><input type="text" name="nice" placeholder="例如：PROTAC, siRNA, 双特异, 海外并购"/></div>
      </div>
      <div class="row">
        <div><label>学历/学校偏好（选填）</label><input type="text" name="edu" placeholder="例如：博士优先；QS200以上；985/211"/></div>
        <div><label>地域/签证等限制（选填）</label><input type="text" name="location" placeholder="例如：上海/苏州；可出差；英文流利"/></div>
      </div>
      <label>补充说明（用于指导AI评估）</label><textarea name="note" placeholder="例如：优先有从PCC→IND推进经验；有license in/out经验；避免频繁跳槽。"></textarea>
    </div>
    <div class="card"><h3>模型与并发</h3>
      <div class="row">
        <div><label>模型名称 <small>(默认 {{model_name}})</small></label><input type="text" name="model_name" value="{{model_name}}"/></div>
        <div><label>每批次并发 <small>(默认 {{max_workers}})</small></label><input type="text" name="workers" value="{{max_workers}}"/></div>
      </div>
      <small>Render 环境变量需配置：MODEL_API_KEY / MODEL_BASE_URL / MODEL_NAME。</small>
    </div>
    <div class="card"><button class="btn" type="submit">开始分析（生成Excel清单）</button>
      <small>提交后会跳到“实时报告”页面，边分析边输出。</small>
    </div>
  </form>
  <div class="card"><h3>历史报告</h3>
    {% if reports %}{% for r in reports %}
      <div class="pill">任务 {{r["id"]}}</div>
      <a href="{{ url_for('view_report', rid=r['id']) }}">查看</a> ·
      <a href="{{ url_for('download_report', rid=r['id']) }}">下载Excel</a>
      <div class="muted">创建：{{r["created_at"]}}；候选数：{{r["counts"]["total"]}}；A+/A：{{r["counts"]["aa"]}}；B/C：{{r["counts"]["bc"]}}</div><br/>
    {% endfor %}{% else %}<div class="muted">暂无</div>{% endif %}
  </div>
  <div class="muted">© Alsos Talent · 合规AI寻访MVP</div>
</div></body></html>"""

STREAM_HTML = """<!DOCTYPE html><html lang="zh"><head><meta charset="utf-8"/>
<meta name="viewport" content="width=device-width,initial-scale=1"/>
<title>实时报告 {{rid}}</title>
<style>
 body{font-family:-apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,Arial;background:#0b0f14;color:#e3e8f2;}
 .wrap{max-width:980px;margin:24px auto;padding:0 16px;}
 .card{background:#121824;border:1px solid #1e2633;border-radius:16px;padding:18px;margin-bottom:16px;}
 .muted{color:#93a1b7;font-size:12px}
 pre{white-space:pre-wrap;word-break:break-word;background:#0b1018;border:1px solid #223044;border-radius:12px;padding:12px;min-height:180px}
 a{color:#7aa0ff;text-decoration:none}
</style></head><body><div class="wrap">
  <div class="card">
    <div>任务 <strong>{{rid}}</strong> · 实时报告 <a href="/">返回</a></div>
    <div class="muted">本页会自动滚动输出；完成后会给出“下载Excel”和“查看榜单”。</div>
  </div>
  <div class="card">
    <pre id="log">🔵 初始化中…</pre>
  </div>
  <script>
    const log = document.getElementById('log');
    const es = new EventSource('/events/{{rid}}');
    function append(t){ log.textContent += '\\n' + t; log.scrollTop = log.scrollHeight; }
    es.onmessage = (e)=>{ append(e.data); }
    es.onerror = ()=>{ append('⚠️ 连接中断，稍后自动重试或手动刷新本页。'); };
  </script>
</div></body></html>"""

RESULTS_HTML = """<!DOCTYPE html><html lang="zh"><head><meta charset="utf-8"/><meta name="viewport" content="width=device-width,initial-scale=1"/>
<title>报告 {{rid}}</title>
<style>
 body{font-family:-apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,Arial; background:#0b0f14; color:#e3e8f2;}
 .wrap{max-width:980px; margin:32px auto; padding:0 16px;}
 .card{background:#121824; border:1px solid #1e2633; border-radius:16px; padding:20px; margin-bottom:18px;}
 h2{margin:0 0 12px;}
 .muted{color:#93a1b7; font-size:12px;}
 table{width:100%; border-collapse:collapse;}
 th,td{border-bottom:1px solid #1f2b3d; padding:10px; text-align:left; vertical-align:top; font-size:14px;}
 .tag{display:inline-block;padding:2px 8px;background:#102033;border:1px solid #223044;border-radius:999px;margin-right:6px;font-size:12px;color:#B8C4D9;}
 a{ color:#7aa0ff; text-decoration:none;}
</style></head><body><div class="wrap">
  <div class="card"><h2>任务 {{rid}}</h2>
    <div class="muted">候选总数：{{counts.total}}；A+/A：{{counts.aa}}；B/C：{{counts.bc}} ·
      <a href="{{ url_for('download_report', rid=rid) }}">下载Excel</a> · <a href="/">返回</a></div>
  </div>
  <div class="card"><h3>重点联系（A+/A）TOP 20</h3>
    <table><thead><tr><th>排名</th><th>姓名/公司</th><th>分数/等级</th><th>摘要</th><th>标签</th></tr></thead><tbody>
      {% for i, row in enumerate(shortlist[:20], start=1) %}
      <tr><td>{{i}}</td>
          <td><strong>{{row.get("name","(未识别)")}}</strong><br/><span class="muted">{{row.get("current_company","")}} · {{row.get("current_title","")}}</span></td>
          <td>{{row.get("overall_score")}} / {{row.get("tier")}}</td>
          <td>{{row.get("fit_summary","")}}</td>
          <td>{% for t in row.get("labels",[]) %}<span class="tag">{{t}}</span>{% endfor %}</td></tr>
      {% endfor %}
    </tbody></table>
  </div>
  <div class="card"><h3>不合适汇总（B/C）示例10条</h3>
    <table><thead><tr><th>姓名</th><th>分数/等级</th><th>主要原因</th><th>备注</th></tr></thead><tbody>
      {% for row in notfit[:10] %}
        <tr><td>{{row.get("name","")}}</td><td>{{row.get("overall_score")}} / {{row.get("tier")}}</td>
            <td>{{", ".join(row.get("risks",[]))}}</td><td>{{row.get("fit_summary","")}}</td></tr>
      {% endfor %}
    </tbody></table>
  </div>
  <div class="muted">© Alsos Talent · 合规AI寻访MVP</div>
</div></body></html>"""

# ------- Helpers -------
def ext_of(name:str)->str:
    name = name.lower()
    for x in ALLOWED_EXT:
        if name.endswith(x):
            return x
    return ""

def read_txt_bytes(b:bytes)->str:
    for enc in ("utf-8","gbk","latin1"):
        try: return b.decode(enc, errors="ignore")
        except Exception: continue
    return ""

def extract_from_pdf(fp)->str:
    if not pdf_extract_text: return ""
    try: return pdf_extract_text(fp) or ""
    except Exception: return ""

def extract_from_docx_bytes(b:bytes)->str:
    if not docx: return ""
    bio = io.BytesIO(b)
    try:
        d = docx.Document(bio)
        return "\n".join(p.text for p in d.paragraphs)
    except Exception:
        return ""

def extract_from_html_bytes(b:bytes)->str:
    if not BeautifulSoup: return read_txt_bytes(b)
    try:
        soup = BeautifulSoup(b, "html.parser")
        for tag in soup(["script","style","noscript"]): tag.extract()
        return soup.get_text("\n", strip=True)
    except Exception:
        return read_txt_bytes(b)

def guess_name(text:str)->str:
    head = (text.strip().splitlines() or [""])[0]
    head = re.sub(r"[\s·•|（）()【】\-_]", " ", head).strip()
    return head[:80]

EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
CN_MOBILE_RE = re.compile(r"(?<!\d)(1[3-9]\d{9})(?!\d)")
GEN_PHONE_RE = re.compile(r"(?<!\d)(\+?\d[\d\s\-()]{6,}\d)(?!\d)")

def extract_contacts(text:str)->Dict[str, Optional[str]]:
    emails = EMAIL_RE.findall(text) or []
    mobiles = CN_MOBILE_RE.findall(text) or []
    phones  = [p.strip() for p in GEN_PHONE_RE.findall(text) if len(p.strip())<=20]
    work_phone = None
    mobile = mobiles[0] if mobiles else None
    for p in phones:
        pp = re.sub(r"\D","",p)
        if mobile and mobile in p: continue
        if len(pp) >= 7:
            work_phone = p; break
    return {"email": (emails[0] if emails else None), "work_phone": work_phone, "mobile": mobile}

YEAR_RE = re.compile(r"(19\d{2}|20\d{2})")
BACHELOR_HINTS = re.compile(r"(本科|学士|Bachelor|B\.Sc|BSc|BS|BA)", re.I)

def estimate_birth_year_str(text:str)->str:
    lines = text.splitlines(); cand_years=[]
    for ln in lines:
        if BACHELOR_HINTS.search(ln):
            cand_years += [int(y) for y in YEAR_RE.findall(ln)]
    if not cand_years: return "不详"
    start_year = min(cand_years); birth = start_year - 18
    return f"约{str(birth)[-2:]}年生"

def minhash_fingerprint(text:str)->str:
    return hashlib.md5(text[:2000].lower().encode("utf-8")).hexdigest()

def truncate(s:str, n:int)->str: return s if len(s)<=n else s[:n]

# ------- LLM call -------
def call_llm(cand_text:str, cand_name:str, job:Dict[str,str])->Dict[str,Any]:
    system = ("You are an expert biotech headhunter assistant. ALWAYS return strict JSON (no markdown). "
              "Scoring: 0-100; Tier: A+,A,B,C (A+/A=strong match). Answer in Chinese.")
    user = {
        "role": job.get("role",""),
        "must_have": job.get("must",""),
        "nice_to_have": job.get("nice",""),
        "min_years": job.get("min_years",""),
        "edu_pref": job.get("edu",""),
        "location_pref": job.get("location",""),
        "note": job.get("note",""),
        "candidate_name": cand_name,
        "candidate_resume": truncate(cand_text, MAX_CHARS_EACH)
    }
    schema_hint = {
      "name":"string","overall_score":"int(0-100)","tier":"one of [A+,A,B,C]",
      "fit_summary":"string (<=120 chars)","risks":["2-4条"],"labels":["关键词"],
      "current_company":"string?","current_title":"string?","location":"string?","remarks":"string"
    }
    prompt = f"""岗位要求与候选文本如下。请输出严格JSON，字段为：
{json.dumps(schema_hint, ensure_ascii=False, indent=2)}

岗位设定:
{json.dumps({k:v for k,v in user.items() if k!='candidate_resume'}, ensure_ascii=False, indent=2)}

候选文本:
{user['candidate_resume']}
"""
    url = MODEL_BASE_URL.rstrip("/") + "/v1/chat/completions"
    headers = {"Authorization": f"Bearer {MODEL_API_KEY}", "Content-Type":"application/json"}
    payload = {"model": MODEL_NAME, "messages":[{"role":"system","content":system},{"role":"user","content":prompt}], "temperature":0.2}
    err=None
    for attempt in range(RETRIES+1):
        try:
            r = requests.post(url, headers=headers, json=payload, timeout=TIMEOUT_SEC); r.raise_for_status()
            content = r.json()["choices"][0]["message"]["content"]
            m = re.search(r"\{.*\}", content, flags=re.S); content = m.group(0) if m else content
            out = json.loads(content); 
            if not out.get("name") and cand_name: out["name"]=cand_name
            return out
        except Exception as e:
            err=e; time.sleep(1.5*(attempt+1))
    return {"name": cand_name or "(未识别)", "overall_score":0, "tier":"C",
            "fit_summary": f"解析失败：{err}", "risks":["LLM调用失败/JSON解析失败"], "labels":[],
            "current_company":"", "current_title":"", "location":"", "remarks": ""}

# ------- Parsing uploads -------
def parse_single_file(name:str, b:bytes)->List[Dict[str,str]]:
    ext = ext_of(name); out=[]
    if ext == ".pdf":
        txt = extract_from_pdf(io.BytesIO(b)) if pdf_extract_text else ""
        out.append({"name":"", "text":txt, "src":name})
    elif ext == ".docx":
        txt = extract_from_docx_bytes(b) if docx else ""
        out.append({"name":"", "text":txt, "src":name})
    elif ext == ".txt":
        out.append({"name":"", "text":read_txt_bytes(b), "src":name})
    elif ext in (".html",".htm"):
        out.append({"name":"", "text":extract_from_html_bytes(b), "src":name})
    elif ext == ".csv":
        # 优先 utf-8，其次 gbk，最后容错
        for enc in ("utf-8", "gbk", None):
            try:
                sio = io.StringIO(b.decode(enc) if enc else b.decode(errors="ignore"))
                reader = csv.DictReader(sio)
                for r in reader:
                    nm = (r.get("Name") or r.get("姓名") or r.get("Candidate") or "").strip()
                    text = " ".join([str(r.get(k,"")) for k in ("Headline","Summary","Experience","Education","Skills","Location")])
                    if not text.strip(): text = " ".join(str(v) for v in r.values())
                    out.append({"name": nm, "text": text.strip(), "src": name})
                break
            except Exception:
                continue
    return out

def parse_uploads(wfs)->List[Dict[str,str]]:
    cands=[]
    for f in wfs:
        if not f.filename: continue
        ext = ext_of(f.filename)
        if not ext: continue
        b = f.read()
        if ext == ".zip":
            try:
                with zipfile.ZipFile(io.BytesIO(b)) as z:
                    for info in z.infolist():
                        if info.is_dir(): continue
                        ext2 = ext_of(info.filename)
                        if not ext2: continue
                        inner = z.read(info.filename)
                        cands.extend(parse_single_file(info.filename, inner))
            except Exception:
                continue
        else:
            cands.extend(parse_single_file(f.filename, b))
    return cands

# ------- Excel output (openpyxl) -------
EXCEL_COLUMNS = [
    "候选人名字","目前所在公司","目前职位","匹配等级（A+/A/B/C）",
    "工作电话","手机","E-mail","年龄预估","目前所在地",
    "契合摘要","风险点","标签","Remarks"
]

def to_excel(rows:List[Dict[str,Any]])->io.BytesIO:
    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active; ws.title = "候选清单"

    ws.append(EXCEL_COLUMNS)
    for r in rows:
        ws.append([
            r.get("name",""),
            r.get("current_company",""),
            r.get("current_title",""),
            r.get("tier",""),
            r.get("work_phone",""),
            r.get("mobile",""),
            r.get("email",""),
            r.get("age_estimate","不详"),
            r.get("location",""),
            r.get("fit_summary",""),
            "，".join(r.get("risks",[]) or []),
            "，".join(r.get("labels",[]) or []),
            r.get("remarks",""),
        ])

    dv = DataValidation(type="list", formula1='"A+,A,B,C"', allow_blank=True,
                        showErrorMessage=True, errorTitle="输入限制", error="请选择 A+ / A / B / C")
    ws.add_data_validation(dv)
    dv.add(f"D2:D5000")

    ws2 = wb.create_sheet("填写说明")
    ws2.append(["字段","说明"])
    instructions = [
        "候选人姓名（中文或英文）",
        "当前就职公司（可从简历/导出文件提取）",
        "当前职位/头衔",
        "从列表选择：A+ / A / B / C",
        "办公电话（无则留空）",
        "手机（无则留空）",
        "邮箱（无则留空）",
        "年龄估算：若识别到“本科入学年份”，计算=入学年-18；否则“不详”",
        "当前所在城市或地区",
        "≤120字，归纳匹配亮点",
        "2–4点主要不匹配/风险",
        "若干关键词，以逗号/顿号分隔（如：ADC, 抗体工程, 临床前, CMC）",
        "长摘要；覆盖现任职责、过往亮点、教育与资质（中文）",
    ]
    for k,v in zip(EXCEL_COLUMNS, instructions):
        ws2.append([k,v])

    bio = io.BytesIO(); wb.save(bio); bio.seek(0); return bio

# ------- Routes -------
@app.route("/", methods=["GET"])
def index():
    items=[{"id":k,"created_at":v.get("created_at"),"counts":v.get("counts",{})} for k,v in REPORTS.items()]
    items.sort(key=lambda x:x["created_at"], reverse=True)
    return render_template_string(INDEX_HTML, reports=items, model_name=MODEL_NAME, max_workers=MAX_WORKERS)

@app.route("/process", methods=["POST"])
def process():
    if not MODEL_API_KEY:
        return "缺少环境变量 MODEL_API_KEY / MODEL_BASE_URL", 400

    global MODEL_NAME  # 需在首次使用前声明

    files = request.files.getlist("files")
    role = request.form.get("role",""); min_years = request.form.get("min_years","")
    must = request.form.get("must",""); nice = request.form.get("nice","")
    edu = request.form.get("edu",""); location = request.form.get("location","")
    note = request.form.get("note","")

    model_name = request.form.get("model_name", MODEL_NAME)
    if model_name:
        MODEL_NAME = model_name

    try:
        workers = int(request.form.get("workers", MAX_WORKERS)); workers = max(1, min(8, workers))
    except Exception:
        workers = MAX_WORKERS

    raw_cands = parse_uploads(files)
    if not raw_cands:
        return "未解析到有效候选文本（请确认ZIP/PDF/HTML/CSV内容）", 400

    pre=[]
    for r in raw_cands:
        text = (r.get("text") or "").strip()
        if not text: continue
        nm = r.get("name") or guess_name(text)
        contacts = extract_contacts(text)
        age_est = estimate_birth_year_str(text)
        pre.append({
            "name": nm, "text": text, "src": r.get("src"),
            "email": contacts.get("email") or "", "work_phone": contacts.get("work_phone") or "", "mobile": contacts.get("mobile") or "",
            "age_estimate": age_est, "fp": minhash_fingerprint(text)
        })

    seen=set(); unique=[]
    for it in pre:
        key=(it["name"], it["fp"])
        if key in seen: continue
        seen.add(key); unique.append(it)

    job={"role":role,"min_years":min_years,"must":must,"nice":nice,"edu":edu,"location":location,"note":note}

    rid = uuid.uuid4().hex[:8]
    q = Queue()
    JOBS[rid] = {"q": q, "done": False, "title": role or "未命名岗位"}

    def runner(items, workers):
        try:
            total = len(items)
            q.put(f"🟢 任务已启动：共 {total} 份候选；并发={workers}")
            results=[]

            if workers <= 1:
                # 顺序处理，最稳
                for idx, it in enumerate(items, start=1):
                    q.put(f"… [{idx}/{total}] 评估：{it['name']}")
                    out = call_llm(it["text"], it["name"], job)
                    row = pack_row(out, it)
                    results.append(row)
                    q.put(f"✅ {row['name']} → {row['tier']} / {row['overall_score']}：{row['fit_summary'][:80]}")
            else:
                from concurrent.futures import ThreadPoolExecutor, as_completed
                q.put("🧵 已启动并发线程…")
                with ThreadPoolExecutor(max_workers=workers) as ex:
                    futs = [ex.submit(lambda x: pack_row(call_llm(x['text'], x['name'], job), x), it) for it in items]
                    done = 0
                    for fu in as_completed(futs):
                        row = fu.result(); results.append(row); done += 1
                        q.put(f"✅ [{done}/{total}] {row['name']} → {row['tier']} / {row['overall_score']}：{row['fit_summary'][:80]}")

            def sort_key(x):
                tier_rank={"A+":0,"A":1,"B":2,"C":3}.get(x.get("tier","C"),3)
                return (tier_rank, -(int(x.get("overall_score") or 0)))

            results_sorted = sorted(results, key=sort_key)
            shortlist=[r for r in results_sorted if r.get("tier") in ("A+","A")]
            notfit=[r for r in results_sorted if r.get("tier") in ("B","C")]

            excel = to_excel(results_sorted)
            REPORTS[rid] = {
                "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "counts": {"total": len(results_sorted), "aa": len(shortlist), "bc": len(notfit)},
                "shortlist": shortlist, "notfit": notfit, "excel": excel
            }
            q.put("🟩 汇总完成：共 {} 人；A+/A：{}；B/C：{}".format(len(results_sorted), len(shortlist), len(notfit)))
            q.put(f"🔗 下载：/download/{rid}    查看：/report/{rid}")
        except Exception as e:
            q.put(f"🟥 任务失败：{e}")
        finally:
            JOBS[rid]["done"] = True
            q.put("[DONE]")

    def pack_row(out, it):
        return {
            "name": out.get("name") or it["name"],
            "overall_score": out.get("overall_score", 0),
            "tier": str(out.get("tier","")).upper(),
            "fit_summary": out.get("fit_summary",""),
            "risks": out.get("risks",[]) or [],
            "labels": out.get("labels",[]) or [],
            "current_company": out.get("current_company",""),
            "current_title": out.get("current_title",""),
            "location": out.get("location",""),
            "remarks": out.get("remarks",""),
            "email": it.get("email",""), "work_phone": it.get("work_phone",""), "mobile": it.get("mobile",""),
            "age_estimate": it.get("age_estimate","不详")
        }

    import threading
    threading.Thread(target=runner, args=(unique, workers), daemon=True).start()
    return render_template_string(STREAM_HTML, rid=rid)

@app.route("/events/<rid>")
def events(rid):
    job = JOBS.get(rid)
    if not job:
        return "任务不存在", 404
    q: Queue = job["q"]

    def gen():
        yield "data: ▶️ 连接已建立\n\n"
        while True:
            try:
                msg = q.get(timeout=12)   # 每 12 秒检测一次
                if msg == "[DONE]":
                    yield "data: 🏁 任务结束\n\n"
                    break
                safe = str(msg).replace("\r", " ").replace("\n", "\\n")
                yield f"data: {safe}\n\n"
            except Empty:
                # 发一个心跳，保持 SSE 通道活跃
                yield f": ping {int(time.time())}\n\n"

    headers = {
        "Content-Type": "text/event-stream",
        "Cache-Control": "no-cache",
        "X-Accel-Buffering": "no",
        "Connection": "keep-alive"
    }
    return Response(gen(), headers=headers)

@app.route("/report/<rid>")
def view_report(rid):
    r = REPORTS.get(rid)
    if not r: return "报告不存在", 404
    return render_template_string(RESULTS_HTML, rid=rid, counts=r["counts"], shortlist=r["shortlist"], notfit=r["notfit"])

@app.route("/download/<rid>")
def download_report(rid):
    r = REPORTS.get(rid)
    if not r: return "报告不存在", 404
    bio = r["excel"]; bio.seek(0)
    return send_file(bio, as_attachment=True, download_name=f"sourcing_report_{rid}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@app.route("/", methods=["HEAD"])
def health():
    return "", 200

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.getenv("PORT", "10000")))
